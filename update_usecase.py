# -*- coding: utf-8 -*-
import docx
import shutil
import sys
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

USECASE_IMG_PATH = r"D:\Java\project\Enterprise AI Assistant\baocao\usecase_cropped.png"

def style_p(p, font_name="Times New Roman", size_pt=13, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.3):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    for r in p.runs:
        r.font.name = font_name
        if size_pt:
            r.font.size = Pt(size_pt)
        if bold is not None:
            r.bold = bold
        if italic is not None:
            r.italic = italic

def set_p_text(p, text, font_name="Times New Roman", size_pt=13, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.3):
    p.text = text
    style_p(p, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, align=align, space_after=space_after, line_spacing=line_spacing)

def process_file(docx_path):
    print("Processing:", docx_path)
    doc = docx.Document(docx_path)
    
    # 1. Update Figure catalog for Hình 2.1
    for p in doc.paragraphs[:65]:
        if "Hình 2.1:" in p.text:
            set_p_text(p, "Hình 2.1: Sơ đồ Use Case tổng quát của hệ thống\t7", font_name="Times New Roman", size_pt=12, bold=False)
            print("Updated Figure catalog entry for Hình 2.1")
            break
            
    # 2. Check and clean duplicate bullet in Section 1.3
    # Look for duplicate P80 / P81
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("- Giám sát và nhật ký kiểm toán: Ghi nhận lịch sử thao tác tải lên"):
            # check if previous paragraph is also nhật ký kiểm toán
            if "Quản trị tổ chức và nhật ký kiểm toán" in doc.paragraphs[i-1].text:
                print("Removing duplicate bullet point at paragraph", i)
                p._p.getparent().remove(p._p)
                break
                
    # 3. Update Actor 2 description to explicitly include 'Chủ bộ phận'
    for p in doc.paragraphs:
        if p.text.strip().startswith("- Trưởng phòng (Department Manager):") or p.text.strip().startswith("- Trưởng phòng / Chủ bộ phận:"):
            set_p_text(p, "- Trưởng phòng / Chủ bộ phận (Department Head): Kế thừa toàn bộ quyền hạn của Nhân viên. Phụ trách quản lý kho tài liệu phòng ban, xét duyệt yêu cầu cấp quyền truy cập, phân quyền thư mục/tệp, quản lý thành viên phòng ban, theo dõi dashboard phân tích và giám sát nhật ký hoạt động nội bộ phòng ban.",
                       font_name="Times New Roman", size_pt=13, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            print("Updated Actor 2 description to match 'Chủ bộ phận'")
            break

    # 4. Locate and replace Hình 2.1
    for i, p in enumerate(doc.paragraphs):
        if "Hình 2.1:" in p.text and i > 65:
            caption_p = p
            img_p = doc.paragraphs[i-1]
            print(f"Found Hình 2.1 at P{i}, replacing image at P{i-1}...")
            
            # Clear old image runs
            img_p.text = ""
            r = img_p.add_run()
            r.add_picture(USECASE_IMG_PATH, width=Inches(6.25))
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_after = Pt(4)
            
            # Update caption
            set_p_text(caption_p, "Hình 2.1: Sơ đồ Use Case tổng quát của hệ thống",
                       font_name="Times New Roman", size_pt=12, bold=True, italic=False,
                       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
            print("Successfully replaced Hình 2.1 image and caption!")
            break

    doc.save(docx_path)
    print("Saved:", docx_path)

# Process both files
file_d = r"D:\Java\project\Enterprise AI Assistant\baocao\baocao_updated.docx"
file_c = r"C:\Users\Admin\Downloads\baocao_updated.docx"
file_bc = r"D:\Java\project\Enterprise AI Assistant\baocao\baocao.docx"

process_file(file_d)
process_file(file_c)
process_file(file_bc)
print("All files updated successfully!")
