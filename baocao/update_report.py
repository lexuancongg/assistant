# -*- coding: utf-8 -*-
"""
Script to update baocao.docx:
1. Rewrite Section I, II, IV to remove AI-sounding clichés, unnatural phrases, and duplicate tokens.
2. Clean all 16 Use Case tables (remove redundant duplicated text, remove placeholders like [CẦN XÁC NHẬN], refine flows).
3. Replace Section V with user's genuine ERD diagram (extracted from ERD.LXC.pdf without watermark) and provide comprehensive, technical entity/relation descriptions based on the 16 real entities.
4. Update Table of Contents (Mục lục) and Figure list (Danh mục hình vẽ) to align with the new structure.
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX_PATH = r"D:\Java\project\Enterprise AI Assistant\baocao\baocao.docx"
BACKUP_PATH = r"D:\Java\project\Enterprise AI Assistant\baocao\baocao_backup_before_edit.docx"
ERD_IMG_PATH = r"D:\Java\project\Enterprise AI Assistant\baocao\erd_final.png"

# Ensure backup exists
shutil.copyfile(DOCX_PATH, BACKUP_PATH)
print("Backup created at:", BACKUP_PATH)

doc = docx.Document(DOCX_PATH)

# Helper function to style paragraph
def style_p(p, font_name="Times New Roman", size_pt=13, bold=False, italic=False, color=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.3):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    for r in p.runs:
        r.font.name = font_name
        if size_pt:
            r.font.size = Pt(size_pt)
        if bold is not None:
            r.bold = bold
        if italic is not None:
            r.italic = italic
        if color:
            r.font.color.rgb = color

def set_p_text(p, text, font_name="Times New Roman", size_pt=13, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.3):
    p.text = text
    style_p(p, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, align=align, space_after=space_after, line_spacing=line_spacing)

# -------------------------------------------------------------
# 1. UPDATE USE CASE TABLES (Table 1 to Table 16)
# -------------------------------------------------------------
use_cases_data = {
    1: {
        "name": "Đăng nhập và Đăng xuất hệ thống (UC-01)",
        "desc": "Xác thực danh tính người dùng bằng tên đăng nhập/email và mật khẩu; cấp phiên làm việc an toàn (JWT/Session token) hoặc thu hồi phiên khi người dùng đăng xuất.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Tên đăng nhập hoặc email.\n- Mật khẩu truy cập.",
        "output": "- Phiên làm việc an toàn (Access Token).\n- Chuyển hướng về giao diện làm việc chính tương ứng với vai trò và quyền hạn.\n- Ghi nhận nhật ký kiểm toán phiên đăng nhập/đăng xuất.",
        "basic": "1. Người dùng nhập tên đăng nhập/email và mật khẩu vào biểu mẫu đăng nhập, nhấn 'Đăng nhập'.\n"
                 "2. Hệ thống kiểm tra định dạng và tính hợp lệ của dữ liệu đầu vào.\n"
                 "3. Hệ thống đối chiếu thông tin xác thực với mật khẩu đã băm trong cơ sở dữ liệu.\n"
                 "4. Hệ thống cấp mã phiên làm việc an toàn, lưu thông tin phiên và ghi nhận sự kiện đăng nhập thành công vào nhật ký kiểm toán (thời gian, địa chỉ IP, mã người dùng).\n"
                 "5. Người dùng được chuyển hướng vào giao diện chính tương ứng với quyền hạn.\n"
                 "6. Khi người dùng nhấn 'Đăng xuất', hệ thống hủy hiệu lực của phiên làm việc, ghi nhận nhật ký đăng xuất và chuyển hướng về trang đăng nhập.",
        "alt": "- 3a. Tên đăng nhập không tồn tại hoặc mật khẩu sai: Hệ thống thông báo 'Tên đăng nhập hoặc mật khẩu không chính xác' và ghi nhận sự kiện đăng nhập thất bại vào nhật ký kiểm toán.\n"
               "- 3b. Tài khoản đang bị khóa: Hệ thống từ chối đăng nhập, thông báo 'Tài khoản của bạn đã bị tạm khóa' và ghi nhận nhật ký cảnh báo bảo mật."
    },
    2: {
        "name": "Quản lý thông tin hồ sơ cá nhân (UC-02)",
        "desc": "Cho phép người dùng xem và cập nhật các thông tin liên lạc cá nhân (họ tên, số điện thoại, ảnh đại diện) trên hệ thống.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Các trường thông tin cập nhật: Họ và tên, Số điện thoại, Ảnh đại diện.",
        "output": "- Dữ liệu hồ sơ được cập nhật trong CSDL.\n- Thông báo cập nhật hồ sơ thành công trên giao diện.",
        "basic": "1. Người dùng chọn mục 'Hồ sơ cá nhân' trên thanh điều hướng.\n"
                 "2. Hệ thống truy vấn CSDL và hiển thị thông tin chi tiết hiện tại của người dùng.\n"
                 "3. Người dùng cập nhật các trường thông tin cần sửa đổi và nhấn 'Lưu thay đổi'.\n"
                 "4. Hệ thống kiểm tra tính hợp lệ của dữ liệu đầu vào (định dạng số điện thoại, hình ảnh).\n"
                 "5. Hệ thống cập nhật bản ghi trong bảng NGUOI_DUNG và hiển thị thông báo thành công.",
        "alt": "- 4a. Số điện thoại hoặc định dạng ảnh đại diện không hợp lệ: Hệ thống hiển thị cảnh báo lỗi tại ô nhập liệu và yêu cầu người dùng kiểm tra lại."
    },
    3: {
        "name": "Đổi mật khẩu tài khoản (UC-03)",
        "desc": "Cung cấp cơ chế an toàn cho phép người dùng chủ động thay đổi mật khẩu đăng nhập định kỳ nhằm bảo vệ tài khoản nội bộ.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Mật khẩu hiện tại.\n- Mật khẩu mới và Xác nhận mật khẩu mới.",
        "output": "- Mật khẩu mới được băm an toàn (BCrypt) và cập nhật vào CSDL.\n- Thông báo đổi mật khẩu thành công và yêu cầu xác thực lại nếu cần.",
        "basic": "1. Người dùng truy cập chức năng 'Đổi mật khẩu' trong phần cài đặt bảo mật.\n"
                 "2. Người dùng nhập mật khẩu hiện tại, mật khẩu mới và xác nhận mật khẩu mới, nhấn 'Cập nhật'.\n"
                 "3. Hệ thống xác thực mật khẩu hiện tại có khớp với dữ liệu đã lưu trữ hay không.\n"
                 "4. Hệ thống kiểm tra độ phức tạp của mật khẩu mới (độ dài tối thiểu 8 ký tự, có chữ hoa, chữ số, ký tự đặc biệt, không trùng mật khẩu cũ).\n"
                 "5. Hệ thống băm mật khẩu mới, lưu vào CSDL, vô hiệu hóa các phiên làm việc cũ và thông báo thành công.",
        "alt": "- 3a. Mật khẩu hiện tại nhập không chính xác: Hệ thống thông báo lỗi và yêu cầu nhập lại.\n"
               "- 4a. Mật khẩu mới không đạt chuẩn bảo mật hoặc không trùng khớp với ô xác nhận: Hệ thống hiển thị thông báo lỗi cụ thể để người dùng sửa đổi."
    },
    4: {
        "name": "Tải lên và tiếp nhận tài liệu doanh nghiệp (UC-04)",
        "desc": "Tiếp nhận tệp tài liệu nghiệp vụ do người dùng tải lên, lưu trữ tệp an toàn, tự động kích hoạt luồng bóc tách văn bản và lập chỉ mục ngữ nghĩa bất đồng bộ phục vụ tra cứu thông minh. Nhân viên được tải tài liệu cá nhân; Trưởng phòng và Quản trị viên được tải tài liệu thuộc phạm vi phòng ban hoặc doanh nghiệp.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Tệp tài liệu nội bộ (hỗ trợ định dạng PDF, DOCX, TXT; dung lượng tối đa 50 MB).\n- Tiêu đề tài liệu, thư mục lưu trữ, phạm vi tài liệu (Cá nhân, Phòng ban, Doanh nghiệp).",
        "output": "- Tệp được lưu trữ an toàn vào kho lưu trữ tệp (Storage/S3).\n- Tạo bản ghi tài liệu (TEP) và bản ghi phiên bản ban đầu (PHIEN_BAN_TEP) ở trạng thái CHO_XU_LY.\n- Thông báo tiếp nhận thành công kèm mã định danh tài liệu.\n- Ghi nhận nhật ký kiểm toán thao tác tải lên.",
        "basic": "1. Người dùng chọn tệp tin văn bản từ thiết bị (PDF, DOCX, TXT, dung lượng <= 50 MB) và chọn thư mục, phạm vi truy cập phù hợp với thẩm quyền.\n"
                 "2. Hệ thống kiểm tra tính hợp lệ của định dạng và dung lượng tệp tin.\n"
                 "3. Hệ thống lưu trữ tệp tin vào kho lưu trữ an toàn, khởi tạo bản ghi tài liệu và phiên bản đầu tiên ở trạng thái Chờ xử lý (CHO_XU_LY).\n"
                 "4. Hệ thống phản hồi thông báo tải lên thành công kèm mã định danh tài liệu.\n"
                 "5. Hệ thống ghi nhận sự kiện tải lên tài liệu vào nhật ký kiểm toán (mã tài liệu, người tải, phòng ban, thời điểm).\n"
                 "6. Hệ thống kích hoạt tác vụ bất đồng bộ sang Tầng Tri thức & AI để khởi chạy luồng bóc tách văn bản, chia đoạn (chunking) và lập chỉ mục vector.\n"
                 "7. Khi tiến trình hoàn tất, hệ thống cập nhật trạng thái tài liệu sang HOAN_THANH (hoặc THAT_BAI nếu xảy ra sự cố) và gửi thông báo cho người dùng.",
        "alt": "- 2a. Tệp tải lên vượt quá dung lượng cho phép hoặc sai định dạng: Hệ thống từ chối tiếp nhận và thông báo lỗi rõ ràng.\n"
               "- 6a. Quá trình bóc tách hoặc tạo vector tại Tầng AI gặp lỗi: Trạng thái tài liệu được cập nhật thành THAT_BAI kèm mã lỗi chi tiết để quản trị viên theo dõi."
    },
    5: {
        "name": "Tìm kiếm và tra cứu tài liệu (UC-05)",
        "desc": "Hỗ trợ người dùng tra cứu nhanh tài liệu nghiệp vụ theo từ khóa, danh mục phòng ban và thời gian ban hành, đảm bảo chỉ hiển thị tài liệu trong phạm vi phân quyền truy cập.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Từ khóa tìm kiếm (tên tài liệu, tóm tắt nội dung).\n- Bộ lọc: Phòng ban, ngày tải lên, định dạng tệp.",
        "output": "- Danh sách tài liệu thỏa mãn điều kiện tìm kiếm và thuộc phạm vi quyền truy cập của người dùng.",
        "basic": "1. Người dùng nhập từ khóa tìm kiếm vào thanh công cụ và chọn bộ lọc nếu cần.\n"
                 "2. Hệ thống xác định danh sách các tài liệu mà người dùng hiện tại có quyền truy cập hợp lệ (theo cá nhân hoặc theo phòng ban).\n"
                 "3. Hệ thống thực hiện tìm kiếm trên tập tài liệu hợp lệ theo tiêu đề, mô tả và nội dung văn bản trích xuất.\n"
                 "4. Hệ thống trả về danh sách kết quả gồm tiêu đề, phiên bản, phòng ban sở hữu và mức quyền của người dùng.",
        "alt": "- 3a. Không tìm thấy tài liệu phù hợp: Hệ thống hiển thị thông báo 'Không tìm thấy tài liệu nào phù hợp với từ khóa' và gợi ý thử lại với từ khóa khác."
    },
    6: {
        "name": "Kiểm soát lịch sử phiên bản tài liệu (UC-06)",
        "desc": "Cho phép cấp quản lý theo dõi toàn bộ vòng đời phát triển của một tài liệu qua các phiên bản cập nhật, bao gồm ngày sửa đổi, người thực hiện và nội dung thay đổi.",
        "actors": "Trưởng phòng, Quản trị viên",
        "input": "- Mã định danh tài liệu cần kiểm tra lịch sử.",
        "output": "- Danh sách toàn bộ các phiên bản của tài liệu với đầy đủ siêu dữ liệu (số phiên bản, thời gian tải, người tải, dung lượng, cờ hiệu lực hiện hành).",
        "basic": "1. Trưởng phòng hoặc Quản trị viên truy cập màn hình chi tiết tài liệu và chọn 'Lịch sử phiên bản'.\n"
                 "2. Hệ thống kiểm tra quyền hạn quản lý của tác nhân đối với tài liệu phòng ban.\n"
                 "3. Hệ thống truy vấn bảng PHIEN_BAN_TEP sắp xếp theo số thứ tự phiên bản giảm dần.\n"
                 "4. Hệ thống hiển thị bảng lịch sử phiên bản trực quan, làm nổi bật phiên bản đang có hiệu lực sử dụng.",
        "alt": "- 2a. Người dùng không có quyền quản lý tài liệu: Hệ thống ẩn chức năng hoặc thông báo từ chối truy cập."
    },
    7: {
        "name": "Khôi phục phiên bản tài liệu (UC-07)",
        "desc": "Cho phép Trưởng phòng hoặc Quản trị viên chỉ định một phiên bản cũ của tài liệu trở thành phiên bản có hiệu lực hiện hành khi phát hiện sai sót trong phiên bản mới.",
        "actors": "Trưởng phòng, Quản trị viên",
        "input": "- Mã tài liệu và số thứ tự phiên bản mục tiêu cần khôi phục.",
        "output": "- Cập nhật phiên bản mục tiêu thành phiên bản hiệu lực hiện hành.\n- Đồng bộ cập nhật chỉ mục tri thức tương ứng tại Tầng AI.\n- Ghi nhận nhật ký kiểm toán thao tác khôi phục phiên bản.",
        "basic": "1. Người dùng có thẩm quyền mở danh sách lịch sử phiên bản của tài liệu.\n"
                 "2. Chọn phiên bản cũ mong muốn và nhấn 'Khôi phục phiên bản'.\n"
                 "3. Hệ thống hiển thị hộp thoại xác nhận kèm cảnh báo về việc chuyển đổi phiên bản hiệu lực.\n"
                 "4. Người dùng xác nhận thực hiện.\n"
                 "5. Hệ thống cập nhật trường ma_phien_ban trong bảng TEP trỏ về phiên bản mục tiêu, đồng bộ chỉ mục ngữ nghĩa tại Tầng AI và ghi nhận sự kiện vào nhật ký kiểm toán.\n"
                 "6. Thông báo khôi phục thành công hiển thị cho người dùng.",
        "alt": "- 4a. Người dùng hủy bỏ thao tác: Hệ thống đóng hộp thoại xác nhận và giữ nguyên trạng thái phiên bản hiện tại."
    },
    8: {
        "name": "Xem trước và tải về tài liệu (UC-08)",
        "desc": "Cung cấp giao diện xem trước nội dung tệp trực tiếp trên trình duyệt web và hỗ trợ tải tệp gốc về thiết bị đối với các tài liệu người dùng có quyền truy cập.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Mã tài liệu, số thứ tự phiên bản và thao tác mong muốn (Xem trước hoặc Tải về).",
        "output": "- Hiển thị nội dung tệp trực tiếp trên trình duyệt hoặc tạo liên kết tải về an toàn có thời hạn (Presigned URL).\n- Ghi nhận nhật ký kiểm toán lượt xem/tải về tài liệu.",
        "basic": "1. Người dùng chọn tài liệu trong danh sách và nhấn 'Xem trước' hoặc 'Tải về'.\n"
                 "2. Hệ thống kiểm tra quyền truy cập tài liệu của người dùng theo phân quyền cá nhân hoặc phòng ban.\n"
                 "3. Nếu có quyền hợp lệ, hệ thống tạo liên kết truy cập an toàn có thời hạn từ kho lưu trữ.\n"
                 "4. Hệ thống ghi nhận sự kiện truy cập/tải về tài liệu vào nhật ký kiểm toán (mã tài liệu, phiên bản, người thực hiện, thời gian).\n"
                 "5. Trình duyệt hiển thị bản xem trước hoặc tự động tải tệp tin về máy tính người dùng.",
        "alt": "- 2a. Người dùng chưa được phân quyền truy cập: Hệ thống hiển thị thông báo 'Bạn chưa có quyền xem tài liệu này' kèm nút 'Gửi yêu cầu cấp quyền'."
    },
    9: {
        "name": "Hỏi đáp tri thức nội bộ bảo mật kèm trích dẫn nguồn (UC-09)",
        "desc": "Tiếp nhận câu hỏi ngôn ngữ tự nhiên từ người dùng, kiểm soát phạm vi phân quyền tài liệu nghiêm ngặt, tìm kiếm các đoạn tri thức tương đồng và tổng hợp câu trả lời chuẩn xác kèm chỉ dẫn nguồn tài liệu gốc (mô hình RAG).",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Câu hỏi bằng ngôn ngữ tự nhiên của người dùng.\n- Mã phiên hội thoại.",
        "output": "- Câu trả lời tổng hợp bằng ngôn ngữ tự nhiên.\n- Danh sách trích dẫn căn cứ gốc (tên tài liệu, số trang, đoạn văn bản trích dẫn).\n- Ghi nhận nhật ký kiểm toán phiên hỏi đáp tri thức.",
        "basic": "1. Người dùng nhập câu hỏi vào ô tương tác trò chuyện và nhấn nút 'Gửi'.\n"
                 "2. Hệ thống tiếp nhận, xác thực phiên làm việc và trích xuất danh sách tài liệu người dùng có quyền đọc hợp lệ.\n"
                 "3. Hệ thống gửi yêu cầu tra cứu ngữ nghĩa sang Tầng Tri thức & AI kèm theo danh mục tài liệu được phép truy cập.\n"
                 "4. Tầng AI chuyển đổi câu hỏi thành vector nhúng, tìm kiếm các phân đoạn văn bản tương đồng cao nhất (trong tập tài liệu hợp lệ) và tổng hợp câu trả lời tự nhiên kèm căn cứ trích dẫn nguồn.\n"
                 "5. Hệ thống lưu tin nhắn người dùng, tin nhắn phản hồi của trợ lý và thông tin trích dẫn vào cơ sở dữ liệu.\n"
                 "6. Hệ thống tự động ghi nhận câu hỏi và các tài liệu nguồn đã sử dụng vào nhật ký kiểm toán để theo dõi tuân thủ an toàn thông tin.\n"
                 "7. Giao diện hiển thị câu trả lời hoàn chỉnh kèm liên kết dẫn chứng đến tài liệu gốc.",
        "alt": "- 2a. Người dùng không có quyền truy cập bất kỳ tài liệu liên quan nào: Hệ thống phản hồi 'Bạn không có quyền truy cập thông tin liên quan đến câu hỏi này' và gợi ý gửi yêu cầu cấp quyền.\n"
               "- 4a. Không tìm thấy nội dung liên quan trong kho tri thức: Hệ thống phản hồi rõ ràng không tìm thấy thông tin phù hợp trong dữ liệu nội bộ."
    },
    10: {
        "name": "Quản lý phiên hội thoại (UC-10)",
        "desc": "Cho phép người dùng tạo các cuộc trò chuyện mới theo từng chủ đề nghiệp vụ, đổi tên tiêu đề phiên, hoặc xóa các phiên hội thoại không còn nhu cầu sử dụng.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Thao tác lựa chọn: Tạo mới phiên, Đổi tên tiêu đề, Xóa phiên.\n- Mã phiên hội thoại.",
        "output": "- Danh sách phiên trò chuyện được cập nhật trên thanh điều hướng bên trái giao diện.\n- Bản ghi cuộc trò chuyện được cập nhật tương ứng trong cơ sở dữ liệu.",
        "basic": "1. Người dùng mở danh sách phiên trò chuyện ở thanh điều hướng bên trái.\n"
                 "2. Chọn thao tác mong muốn: nhấn 'Cuộc trò chuyện mới', 'Đổi tên' hoặc 'Xóa phiên'.\n"
                 "3. Nếu chọn 'Đổi tên', người dùng nhập tiêu đề mới và hệ thống lưu thay đổi vào bảng PHIEN_CHAT.\n"
                 "4. Nếu chọn 'Xóa phiên', hệ thống yêu cầu xác nhận và tiến hành xóa phiên hội thoại cùng các tin nhắn liên quan.\n"
                 "5. Giao diện cập nhật lại danh sách phiên trò chuyện theo trạng thái mới.",
        "alt": "- 3a. Phiên không thuộc quyền sở hữu của người dùng: Hệ thống từ chối thao tác và thông báo không có quyền."
    },
    11: {
        "name": "Đánh giá phản hồi câu trả lời của trợ lý (UC-11)",
        "desc": "Thu thập phản hồi đánh giá của người dùng bao gồm bình chọn Hài lòng hoặc Không hài lòng và nhận xét góp ý chi tiết đối với từng câu trả lời do trợ lý đưa ra, làm căn cứ cải thiện chất lượng tri thức.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Mã tin nhắn cần đánh giá.\n- Đánh giá Hài lòng hoặc Không hài lòng.\n- Nhận xét chi tiết và lý do cụ thể (tùy chọn).",
        "output": "- Bản ghi đánh giá mới được lưu trữ trong bảng DANH_GIA_TIN_NHAN trong CSDL.\n- Thông báo ghi nhận phản hồi thành công hiển thị cho người dùng.",
        "basic": "1. Dưới mỗi câu trả lời của trợ lý, người dùng nhấn vào biểu tượng Hài lòng hoặc Không hài lòng.\n"
                 "2. Nếu chọn Không hài lòng, hệ thống hiển thị biểu mẫu nhập nhận xét chi tiết để người dùng nêu rõ lý do góp ý.\n"
                 "3. Người dùng nhấn nút 'Gửi phản hồi'.\n"
                 "4. Hệ thống kiểm tra dữ liệu và lưu bản ghi đánh giá vào bảng DANH_GIA_TIN_NHAN trong cơ sở dữ liệu.\n"
                 "5. Hệ thống hiển thị thông báo cảm ơn và ghi nhận đánh giá thành công.",
        "alt": "- 4a. Lỗi kết nối hệ thống: Hệ thống thông báo 'Không thể ghi nhận phản hồi vào lúc này, vui lòng thử lại sau'."
    },
    12: {
        "name": "Trích xuất và tóm tắt văn bản thông minh (UC-12)",
        "desc": "Tự động phân tích các tài liệu dài hoặc phức tạp, cô đọng nội dung thành bản tóm tắt ngắn gọn có cấu trúc (luận điểm chính, lưu ý quan trọng) phục vụ đọc nhanh cho nhân viên.",
        "actors": "Nhân viên, Trưởng phòng, Quản trị viên",
        "input": "- Mã tài liệu cần tóm tắt.\n- Tùy chọn định dạng tóm tắt: Tóm tắt ngắn gọn, Tóm tắt theo đề mục, Bản ghi nhớ nội dung chính.",
        "output": "- Văn bản tóm tắt có cấu trúc rõ ràng hiển thị trực tiếp trên giao diện và hỗ trợ sao chép nhanh.",
        "basic": "1. Người dùng chọn một tài liệu trong kho và nhấn nút 'Tóm tắt nội dung'.\n"
                 "2. Hệ thống kiểm tra quyền truy cập tài liệu của người dùng.\n"
                 "3. Hệ thống lấy toàn bộ văn bản của phiên bản tài liệu hiện hành và chuyển sang Tầng Tri thức & AI.\n"
                 "4. Tầng AI phân tích cấu trúc, xác định các ý trọng tâm và tổng hợp thành bản tóm tắt súc tích theo định dạng được yêu cầu.\n"
                 "5. Hệ thống trả về văn bản tóm tắt trên màn hình kèm các nút 'Sao chép' hoặc 'Xuất ra tệp'.",
        "alt": "- 2a. Người dùng không có quyền truy cập tài liệu: Hệ thống từ chối yêu cầu và thông báo thiếu thẩm quyền."
    },
    13: {
        "name": "Gửi yêu cầu cấp quyền tài liệu (UC-13)",
        "desc": "Cho phép người dùng gửi phiếu đề nghị xin cấp quyền truy cập đối với các tài liệu nghiệp vụ phòng ban khác khi có nhu cầu công việc. Hỗ trợ 3 mức quyền: XEM, CHỈNH SỬA, QUẢN TRỊ. Phiếu của Nhân viên gửi tới Trưởng phòng duyệt; phiếu của Trưởng phòng gửi lên Quản trị viên duyệt.",
        "actors": "Nhân viên, Trưởng phòng",
        "input": "- Mã tài liệu cần xin quyền.\n- Mức quyền mong muốn (XEM, CHỈNH SỬA hoặc QUẢN TRỊ).\n- Lý do công việc cụ thể.",
        "output": "- Bản ghi yêu cầu mới được tạo ở trạng thái Chờ duyệt (CHO_DUYET) trong bảng YEU_CAU_CAP_QUYEN.\n- Thông báo gửi yêu cầu thành công hiển thị cho người xin và thông báo tự động chuyển đến cấp có thẩm quyền phê duyệt.",
        "basic": "1. Người dùng tìm kiếm thấy tài liệu nhưng bị giới hạn quyền truy cập, nhấn vào nút 'Gửi yêu cầu cấp quyền'.\n"
                 "2. Hệ thống hiển thị biểu mẫu yêu cầu cấp quyền gồm thông tin tài liệu, lựa chọn mức quyền (XEM, CHỈNH SỬA, QUẢN TRỊ) và ô nhập lý do công việc.\n"
                 "3. Người dùng lựa chọn mức quyền, điền lý do chi tiết và nhấn nút 'Gửi yêu cầu'.\n"
                 "4. Hệ thống kiểm tra tính hợp lệ của lý do và xác định thẩm quyền phê duyệt: nếu người xin là Nhân viên thì chỉ định Trưởng phòng duyệt; nếu người xin là Trưởng phòng thì chỉ định Quản trị viên duyệt.\n"
                 "5. Hệ thống lưu bản ghi yêu cầu vào cơ sở dữ liệu với trạng thái CHO_DUYET và tạo thông báo hệ thống gửi đến cấp phê duyệt.\n"
                 "6. Hệ thống hiển thị thông báo gửi phiếu yêu cầu thành công.",
        "alt": "- 3a. Người dùng bỏ trống trường lý do: Hệ thống yêu cầu điền đầy đủ lý do trước khi gửi."
    },
    14: {
        "name": "Phê duyệt yêu cầu cấp quyền tài liệu (UC-14)",
        "desc": "Cung cấp giao diện quản lý cho Trưởng phòng hoặc Quản trị viên xem xét các phiếu xin cấp quyền tài liệu, đưa ra quyết định Phê duyệt (DA_PHE_DUYET) hoặc Từ chối (TU_CHOI) kèm theo lý do cụ thể.",
        "actors": "Trưởng phòng, Quản trị viên",
        "input": "- Mã phiếu yêu cầu.\n- Quyết định: Phê duyệt (DA_PHE_DUYET) hoặc Từ chối (TU_CHOI).\n- Ghi chú hoặc Lý do từ chối (bắt buộc khi từ chối).",
        "output": "- Trạng thái phiếu yêu cầu được cập nhật thành DA_PHE_DUYET hoặc TU_CHOI.\n- Nếu được duyệt, hệ thống tự động tạo bản ghi phân quyền tương ứng trong bảng QUYEN_TEP hoặc QUYEN_THU_MUC.\n- Tạo thông báo hệ thống gửi phản hồi kết quả cho người yêu cầu.\n- Ghi nhận nhật ký kiểm toán thao tác phê duyệt/từ chối.",
        "basic": "1. Người có thẩm quyền (Trưởng phòng/Quản trị viên) mở danh mục 'Phê duyệt yêu cầu cấp quyền'.\n"
                 "2. Hệ thống hiển thị danh sách các phiếu yêu cầu đang ở trạng thái CHO_DUYET thuộc phạm vi quản lý.\n"
                 "3. Người duyệt chọn một phiếu yêu cầu để xem chi tiết: người yêu cầu, phòng ban, mức quyền xin cấp, lý do.\n"
                 "4. Người duyệt đưa ra quyết định: nhấn 'Phê duyệt' hoặc 'Từ chối' (nhập lý do từ chối).\n"
                 "5. Nếu phê duyệt, hệ thống cập nhật trạng thái phiếu thành DA_PHE_DUYET, tự động tạo bản ghi phân quyền tương ứng trong CSDL và gửi thông báo cho người xin.\n"
                 "6. Nếu từ chối, hệ thống cập nhật trạng thái phiếu thành TU_CHOI kèm lý do và gửi thông báo phản hồi.\n"
                 "7. Hệ thống ghi nhận sự kiện phê duyệt hoặc từ chối vào nhật ký kiểm toán và làm mới danh sách yêu cầu.",
        "alt": "- 4a. Yêu cầu đã được xử lý bởi người khác trước đó: Hệ thống cảnh báo trạng thái đã thay đổi và làm mới danh sách."
    },
    15: {
        "name": "Quản trị danh mục người dùng và phòng ban (UC-15)",
        "desc": "Cho phép Quản trị viên hệ thống thiết lập cơ cấu tổ chức, quản lý danh sách phòng ban, tạo mới, chỉnh sửa, gán vai trò hoặc khóa tài khoản người dùng trên toàn doanh nghiệp.",
        "actors": "Quản trị viên",
        "input": "- Dữ liệu phòng ban: Tên phòng ban, mã công ty.\n- Dữ liệu người dùng: Họ và tên, email, mật khẩu khởi tạo, phòng ban trực thuộc, vai trò (NHAN_VIEN, TRUONG_PHONG, ADMIN), trạng thái tài khoản.",
        "output": "- Các bản ghi được tạo mới hoặc cập nhật tương ứng trong cơ sở dữ liệu (bảng PHONG_BAN, NGUOI_DUNG).\n- Thông báo thao tác thành công hiển thị cho Quản trị viên.\n- Ghi nhận nhật ký kiểm toán thao tác quản trị người dùng/phòng ban.",
        "basic": "1. Quản trị viên truy cập phân hệ 'Quản trị hệ thống' -> 'Người dùng & Phòng ban'.\n"
                 "2. Lựa chọn thao tác: Tạo mới phòng ban, Thêm tài khoản người dùng, Gán vai trò hoặc Khóa tài khoản.\n"
                 "3. Quản trị viên nhập thông tin biểu mẫu và nhấn nút 'Lưu'.\n"
                 "4. Hệ thống kiểm tra các ràng buộc: tính duy nhất của email, định dạng email, phòng ban trực thuộc hợp lệ.\n"
                 "5. Hệ thống thực hiện cập nhật CSDL và ghi nhận sự kiện vào bảng NHAT_KY_HOAT_DONG.\n"
                 "6. Hệ thống cập nhật bảng dữ liệu quản trị trên màn hình và thông báo thành công.",
        "alt": "- 4a. Email đã tồn tại trong hệ thống: Báo lỗi trùng lặp dữ liệu và yêu cầu kiểm tra lại thông tin."
    },
    16: {
        "name": "Giám sát nhật ký kiểm toán hệ thống (UC-16)",
        "desc": "Cung cấp cho Quản trị viên khả năng tra cứu, lọc và theo dõi toàn bộ các sự kiện vận hành nghiệp vụ quan trọng (đăng nhập, tải lên, xóa tài liệu, phê duyệt quyền, truy vấn tri thức) phục vụ an toàn thông tin.",
        "actors": "Quản trị viên",
        "input": "- Bộ lọc kiểm toán: Khoảng thời gian, Người thực hiện, Loại hành động (DANG_NHAP, TAI_LEN, TRUY_VAN, KHOI_PHUC, PHE_DUYET,...), Mã tài liệu liên quan.",
        "output": "- Bảng danh sách nhật ký chi tiết (Dấu thời gian, Mã nhân viên, Loại hành động, Mã tài liệu, Chi tiết thay đổi) kèm phân trang tra cứu.",
        "basic": "1. Quản trị viên chọn chức năng 'Nhật ký kiểm toán' trên thanh điều hướng quản trị.\n"
                 "2. Thiết lập các tiêu chí lọc sự kiện theo nhu cầu điều tra an ninh.\n"
                 "3. Nhấn 'Truy vấn nhật ký'.\n"
                 "4. Hệ thống kiểm tra thẩm quyền Quản trị viên.\n"
                 "5. Hệ thống truy vấn bảng NHAT_KY_HOAT_DONG theo điều kiện lọc và trả về danh sách sự kiện.\n"
                 "6. Hệ thống hiển thị bảng nhật ký chi tiết có phân trang và hỗ trợ tra cứu nhanh.",
        "alt": "- 4a. Tác nhân không đủ thẩm quyền Quản trị viên: Hệ thống chặn truy cập và ghi cảnh báo cố gắng truy cập trái phép."
    }
}

for uc_id, uc_info in use_cases_data.items():
    t = doc.tables[uc_id]
    field_map = [
        (0, "Use-case name", uc_info["name"]),
        (1, "Description", uc_info["desc"]),
        (2, "Actors", uc_info["actors"]),
        (3, "Input", uc_info["input"]),
        (4, "Output", uc_info["output"]),
        (5, "Basic flow", uc_info["basic"]),
        (6, "Alternative flow", uc_info["alt"]),
    ]
    for row_idx, label, val in field_map:
        row = t.rows[row_idx]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.text = label
        c1.text = val
        
        # style c0
        p0 = c0.paragraphs[0]
        style_p(p0, font_name="Times New Roman", size_pt=12, bold=True, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2, line_spacing=1.2)
        
        # style c1
        for p in c1.paragraphs:
            style_p(p, font_name="Times New Roman", size_pt=11.5, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2, line_spacing=1.25)

print("Updated 16 Use Case tables successfully!")

# -------------------------------------------------------------
# 2. UPDATE TEXT IN SECTION I, II, IV
# -------------------------------------------------------------
# P64: 1.1 Context
set_p_text(doc.paragraphs[64], 
           "Trong quá trình vận hành của doanh nghiệp, khối lượng tài liệu nội bộ—bao gồm quy chế, quy trình nghiệp vụ, tài liệu kỹ thuật, chính sách nhân sự và biểu mẫu hành chính—không ngừng gia tăng theo quy mô tổ chức. Phần lớn nguồn tài liệu này hiện đang được lưu trữ phân tán dưới nhiều định dạng khác nhau (PDF, Word, bảng tính, văn bản thuần) trên máy tính cá nhân của nhân viên, qua các ứng dụng trao đổi tin nhắn hoặc trong các thư mục chia sẻ nội bộ. Thực tế này gây ra nhiều khó khăn trong quản lý và vận hành:")

# P65: Time
set_p_text(doc.paragraphs[65],
           "- Mất nhiều thời gian tra cứu: Nhân viên gặp trở ngại khi cần tìm kiếm chính xác quy định, hướng dẫn nghiệp vụ hoặc biểu mẫu phục vụ công việc, làm giảm hiệu suất lao động.")

# P66: Versioning
set_p_text(doc.paragraphs[66],
           "- Khó kiểm soát phiên bản và thiếu đồng bộ: Việc cập nhật tài liệu qua nhiều lần chỉnh sửa dễ dẫn đến tình trạng áp dụng nhầm các phiên bản cũ đã hết hiệu lực thi hành.")

# P67: Leak
set_p_text(doc.paragraphs[67],
           "- Nguy cơ thất thoát và rò rỉ thông tin: Cơ chế chia sẻ tệp thủ công không có sự kiểm soát quyền truy cập chặt chẽ, tiềm ẩn rủi ro lộ lọt dữ liệu nội bộ giữa các phòng ban.")

# P68: Support overload
set_p_text(doc.paragraphs[68],
           "- Quá tải nhân sự hỗ trợ: Các phòng ban chức năng (Hành chính - Nhân sự, Kế toán, Kỹ thuật) thường xuyên phải dành nhiều thời gian để giải đáp lặp đi lặp lại những câu hỏi nghiệp vụ mang tính định kỳ.")

# P69: Heading 1.2
set_p_text(doc.paragraphs[69], "1.2. Mục tiêu và phạm vi xây dựng hệ thống", bold=True, space_after=4)

# P70: Objective
set_p_text(doc.paragraphs[70],
           "Đề tài tập trung nghiên cứu, phân tích và thiết kế hệ thống website quản lý và khai thác dữ liệu cho doanh nghiệp, có tích hợp trợ lý AI thông minh. Mục tiêu hướng tới là xây dựng một nền tảng tập trung hỗ trợ quản trị toàn bộ vòng đời tài liệu, tổ chức cây thư mục khoa học, phân quyền truy cập đa cấp theo vai trò và phòng ban, đồng thời cung cấp khả năng tra cứu ngữ nghĩa và hỏi đáp thông minh dựa trên kho tài liệu nội bộ (áp dụng mô hình RAG - Retrieval-Augmented Generation).")

# P71: Scope
set_p_text(doc.paragraphs[71],
           "Báo cáo chuyên đề này tập trung vào giai đoạn phân tích và thiết kế hệ thống ở mức logic: khảo sát nghiệp vụ thực tế, mô hình hóa các Use Case và kịch bản thực thi, xây dựng sơ đồ tuần tự thể hiện các luồng xử lý động, thiết kế kiến trúc phân tầng logic và thiết kế cơ sở dữ liệu quan hệ (ERD) hoàn chỉnh.")

# P72: Heading 1.3
set_p_text(doc.paragraphs[72], "1.3. Yêu cầu chức năng", bold=True, space_after=4)

# P73 - P80: Functional requirements
set_p_text(doc.paragraphs[73],
           "- Quản lý định danh và phiên làm việc: Xác thực người dùng (đăng nhập, đăng xuất, đổi mật khẩu), cấp và kiểm soát phiên làm việc an toàn, quản lý thông tin hồ sơ cá nhân.")
set_p_text(doc.paragraphs[74],
           "- Quản lý kho tài liệu và cây thư mục: Hỗ trợ tạo lập, phân loại cây thư mục; tiếp nhận tải lên các tệp tài liệu đa định dạng (PDF, DOCX, TXT); theo dõi lịch sử các phiên bản và cho phép khôi phục về phiên bản trước khi cần thiết.")
set_p_text(doc.paragraphs[75],
           "- Phân quyền tài liệu và duyệt cấp quyền: Thiết lập quyền truy cập theo 3 cấp độ (XEM, CHỈNH SỬA, QUẢN TRỊ) áp dụng linh hoạt cho cá nhân hoặc phòng ban; cung cấp quy trình gửi yêu cầu xin cấp quyền và giao diện phê duyệt/từ chối cho cấp quản lý.")
set_p_text(doc.paragraphs[76],
           "- Trích xuất và lập chỉ mục tri thức: Tự động trích xuất nội dung văn bản từ tài liệu tải lên, phân đoạn văn bản (chunking) và tạo nhúng vector ngữ nghĩa để phục vụ tìm kiếm thông minh.")
set_p_text(doc.paragraphs[77],
           "- Hỏi đáp tri thức kèm trích dẫn nguồn: Cho phép người dùng đặt câu hỏi bằng ngôn ngữ tự nhiên; hệ thống tự động lọc tài liệu theo đúng quyền truy cập của người dùng, tìm kiếm nội dung liên quan và tổng hợp câu trả lời chính xác kèm liên kết dẫn chứng đến văn bản gốc.")
set_p_text(doc.paragraphs[78],
           "- Quản lý phiên hội thoại và đánh giá: Lưu trữ lịch sử hỏi đáp theo từng phiên làm việc; hỗ trợ người dùng gửi đánh giá hài lòng/không hài lòng và góp ý chi tiết đối với câu trả lời của trợ lý.")
set_p_text(doc.paragraphs[79],
           "- Quản trị tổ chức và nhật ký kiểm toán: Cho phép Quản trị viên quản lý danh mục phòng ban, danh sách nhân viên, gán vai trò; ghi nhận và tra cứu nhật ký hoạt động (Audit Log) đối với toàn bộ các thao tác trọng yếu trên hệ thống.")

# P82 - P86: Section 2.1 Actors
set_p_text(doc.paragraphs[82],
           "Hệ thống xác định 3 tác nhân nghiệp vụ tương tác trực tiếp gồm: Nhân viên, Trưởng phòng và Quản trị viên. Giữa các tác nhân được áp dụng quan hệ khái quát hóa (Generalization) trong UML nhằm thiết lập cấu trúc phân quyền kế thừa chặt chẽ và loại bỏ sự trùng lặp trong kiểm tra logic:")
set_p_text(doc.paragraphs[83],
           "- Nhân viên (Employee): Tác nhân cơ sở trong hệ thống. Nhân viên thực hiện đăng nhập, cập nhật hồ sơ cá nhân, xem và tải tài liệu được phân quyền, tương tác hỏi đáp với trợ lý AI, đánh giá chất lượng câu trả lời và gửi yêu cầu xin cấp quyền khi cần truy cập tài liệu ngoài phạm vi được giao.")
set_p_text(doc.paragraphs[84],
           "- Trưởng phòng (Department Manager): Kế thừa toàn bộ quyền hạn của Nhân viên. Trưởng phòng quản lý kho tài liệu thuộc phạm vi phòng ban mình phụ trách, xem xét và phê duyệt hoặc từ chối các yêu cầu xin cấp quyền tài liệu của nhân viên, đồng thời có quyền khôi phục các phiên bản tài liệu cũ khi phát hiện sai sót.")
set_p_text(doc.paragraphs[85],
           "- Quản trị viên (Admin): Kế thừa toàn bộ quyền hạn của Trưởng phòng và Nhân viên. Quản trị viên phụ trách quản lý danh mục phòng ban, tài khoản người dùng, cấu hình quyền hạn toàn hệ thống và theo dõi nhật ký kiểm toán hoạt động để đảm bảo an toàn thông tin.")
set_p_text(doc.paragraphs[86],
           "Cấu trúc kế thừa giữa các tác nhân giúp chính sách phân quyền luôn nhất quán và dễ dàng mở rộng khi phát triển thêm các vai trò nghiệp vụ mới.")

# P92 - P93: UC-09 includes / extends
set_p_text(doc.paragraphs[92],
           "Quan hệ «include» của UC-09: Khi thực hiện chức năng hỏi đáp tri thức, hệ thống bắt buộc tự động thực hiện các luồng con: xác thực phạm vi phân quyền tài liệu của người dùng, tìm kiếm ngữ nghĩa theo nội dung câu hỏi và lưu vết lịch sử tương tác vào cơ sở dữ liệu.")
set_p_text(doc.paragraphs[93],
           "Quan hệ «extend» của UC-09: Trong trường hợp câu hỏi của người dùng nằm ngoài phạm vi tài liệu được phân quyền hoặc không tìm thấy thông tin phù hợp trong kho tri thức, Use Case cảnh báo từ chối truy cập sẽ được kích hoạt để phản hồi phù hợp.")

# P166 - P170: Section IV 4 Layer Architecture
set_p_text(doc.paragraphs[166],
           "Kiến trúc hệ thống được thiết kế theo mô hình 4 tầng logic độc lập, bảo đảm nguyên tắc phân tách trách nhiệm (Separation of Concerns) và khả năng mở rộng linh hoạt:")
set_p_text(doc.paragraphs[167],
           "- Tầng Giao diện (Presentation Layer): Cung cấp giao diện tương tác trực quan cho người dùng và quản trị viên trên nền tảng web, bao gồm màn hình quản lý tài liệu, cây thư mục, khung chat hỏi đáp trợ lý AI, biểu mẫu xin cấp quyền và bảng điều khiển quản trị.")
set_p_text(doc.paragraphs[168],
           "- Tầng Nghiệp vụ (Business Logic Layer): Xử lý toàn bộ logic nghiệp vụ cốt lõi của hệ thống, bao gồm xác thực định danh (Authentication), kiểm soát phân quyền đa cấp (Authorization), quản lý vòng đời tài liệu, điều phối quy trình phê duyệt yêu cầu và quản lý phiên hội thoại.")
set_p_text(doc.paragraphs[169],
           "- Tầng Tri thức & AI (Knowledge & AI Processing Layer): Đảm nhận nhiệm vụ bóc tách nội dung văn bản (Document Parser), phân đoạn tài liệu (Chunking), tạo nhúng vector ngữ nghĩa (Embedding Model), tìm kiếm tương đồng và sinh câu trả lời tự nhiên kèm trích dẫn nguồn (LLM/RAG Pipeline).")
set_p_text(doc.paragraphs[170],
           "- Tầng Dữ liệu & Lưu trữ (Data Access & Persistence Layer): Chịu trách nhiệm lưu trữ bền vững và bảo đảm toàn vẹn dữ liệu, bao gồm cơ sở dữ liệu quan hệ (PostgreSQL) lưu trữ thông tin thực thể nghiệp vụ, cơ sở dữ liệu vector lưu trữ các vector nhúng ngữ nghĩa, và dịch vụ lưu trữ tệp tin (Object Storage/S3) lưu trữ các tệp nhị phân gốc.")

print("Updated text in Section I, II, IV successfully!")

# -------------------------------------------------------------
# 3. REWRITE SECTION V: THIẾT KẾ CƠ SỞ DỮ LIỆU MỨC LOGIC (ERD)
# -------------------------------------------------------------
# Let's see: in doc.paragraphs, Section V begins at P173.
# Let's inspect paragraphs 173 to end:
# P173: V. THIẾT KẾ CƠ SỞ DỮ LIỆU MỨC LOGIC (ERD)
# P174: 5.1. ...
# ... to end.

# First, remove old paragraphs from 174 onwards, keeping 173
total_p = len(doc.paragraphs)
print(f"Paragraphs before Section V rewrite: {total_p}")

# To safely replace Section V:
# We keep P173 (Heading V), and we can rewrite/replace the rest of paragraphs.
# In python-docx, we can remove paragraphs by removing p._p from parent.
for i in range(total_p - 1, 173, -1):
    p = doc.paragraphs[i]
    p._p.getparent().remove(p._p)

print("Removed old Section V content. Remaining paragraphs:", len(doc.paragraphs))

# Now let's add the new comprehensive Section V!
p_v_title = doc.paragraphs[173]
set_p_text(p_v_title, "V. THIẾT KẾ CƠ SỞ DỮ LIỆU MỨC LOGIC (ERD)", bold=True, size_pt=14, space_after=8)

# 5.1. Sơ đồ quan hệ thực thể (ERD) tổng thể hệ thống
p_51 = doc.add_paragraph()
set_p_text(p_51, "5.1. Sơ đồ quan hệ thực thể (ERD) tổng thể hệ thống", bold=True, size_pt=13, space_after=6)

p_intro = doc.add_paragraph()
set_p_text(p_intro, 
           "Cơ sở dữ liệu mức logic của hệ thống được thiết kế chuẩn hóa nhằm đáp ứng toàn diện các yêu cầu nghiệp vụ về quản trị tổ chức, phân quyền tài liệu đa cấp, quy trình phê duyệt, bóc tách tri thức phục vụ mô hình RAG và lưu vết kiểm toán. Sơ đồ ERD tổng thể bao gồm 16 thực thể quan hệ, được tổ chức thành 4 phân hệ chức năng đồng bộ:")

# Insert the high-res ERD diagram
p_img = doc.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img.paragraph_format.space_after = Pt(4)
r_img = p_img.add_run()
r_img.add_picture(ERD_IMG_PATH, width=Inches(6.25))

# Image caption
p_cap = doc.add_paragraph()
set_p_text(p_cap, "Hình 5.1: Sơ đồ quan hệ thực thể (ERD) mức logic của hệ thống", bold=True, size_pt=12, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# 5.2. Đặc tả chi tiết các thực thể và ràng buộc dữ liệu
p_52 = doc.add_paragraph()
set_p_text(p_52, "5.2. Đặc tả chi tiết các thực thể và ràng buộc dữ liệu", bold=True, size_pt=13, space_after=6)

p_desc_intro = doc.add_paragraph()
set_p_text(p_desc_intro,
           "Dưới đây là đặc tả chi tiết về cấu trúc thuộc tính, kiểu dữ liệu, các ràng buộc khóa chính (PK), khóa ngoại (FK) và mối liên kết logic giữa 16 thực thể trong mô hình:")

# Subsystem 1: Tổ chức & Người dùng
p_sub1 = doc.add_paragraph()
set_p_text(p_sub1, "5.2.1. Phân hệ Quản trị Tổ chức và Người dùng", bold=True, size_pt=12.5, space_after=4)

p_sub1_desc = doc.add_paragraph()
set_p_text(p_sub1_desc,
           "Phân hệ quản lý cơ cấu tổ chức doanh nghiệp và danh sách nhân sự tham gia vận hành hệ thống, gồm 3 thực thể:\n"
           "1. Bảng Cong_ty (Doanh nghiệp):\n"
           "   - ma_cty (BIGINT, PK): Mã định danh duy nhất của doanh nghiệp.\n"
           "   - ten (VARCHAR(255)): Tên đầy đủ của doanh nghiệp.\n"
           "   - email (VARCHAR(255)): Địa chỉ email liên hệ chính thức.\n"
           "   - sđt (VARCHAR(50)): Số điện thoại liên hệ.\n"
           "   - dia_chi (TEXT): Địa chỉ trụ sở công ty.\n"
           "   - han_muc (INTEGER): Hạn mức tài nguyên được cấp (dung lượng lưu trữ tối đa hoặc số lượng tài khoản).\n"
           "   - thoi_gian_tao (TIMESTAMP): Thời điểm khởi tạo bản ghi doanh nghiệp.\n"
           "\n"
           "2. Bảng Phong_ban (Phòng ban trực thuộc):\n"
           "   - ma_phong_ban (BIGINT, PK): Mã định danh phòng ban.\n"
           "   - ma_cty (BIGINT, FK -> Cong_ty.ma_cty): Xác định phòng ban trực thuộc doanh nghiệp nào (quan hệ 1-N: một công ty có nhiều phòng ban).\n"
           "   - ten_pb (VARCHAR(255)): Tên phòng ban (ví dụ: Phòng Kế toán, Phòng Nhân sự, Phòng Kỹ thuật).\n"
           "   - thoi_gian_thanh_lap (TIMESTAMP): Thời điểm thành lập phòng ban.\n"
           "\n"
           "3. Bảng Nguoi_dung (Tài khoản nhân sự):\n"
           "   - ma_nv (BIGINT, PK): Mã định danh nhân viên.\n"
           "   - thuoc_cty (BIGINT, FK -> Cong_ty.ma_cty): Xác định nhân viên trực thuộc công ty.\n"
           "   - thuoc_phong_ban (BIGINT, FK -> Phong_ban.ma_phong_ban): Phòng ban mà nhân viên đang công tác (quan hệ 1-N: một phòng ban có nhiều nhân viên).\n"
           "   - ho_va_ten (VARCHAR(255)): Họ và tên đầy đủ của nhân viên.\n"
           "   - email (VARCHAR(255), UNIQUE): Email định danh đăng nhập vào hệ thống.\n"
           "   - mat_khau (TEXT): Mật khẩu truy cập đã được mã hóa băm bảo mật (BCrypt).\n"
           "   - vai_tro (VARCHAR(20)): Vai trò quyền hạn trong hệ thống, gồm: NHAN_VIEN, TRUONG_PHONG, ADMIN.\n"
           "   - trang_thai (VARCHAR(20)): Trạng thái tài khoản (DANG_HOAT_DONG, BI_KHOA).\n"
           "   - thoi_gian_tao (TIMESTAMP): Thời điểm tạo tài khoản.")

# Subsystem 2: Tài liệu & Phân quyền
p_sub2 = doc.add_paragraph()
set_p_text(p_sub2, "5.2.2. Phân hệ Quản lý Kho tài liệu và Phân quyền truy cập", bold=True, size_pt=12.5, space_after=4)

p_sub2_desc = doc.add_paragraph()
set_p_text(p_sub2_desc,
           "Phân hệ đóng vai trò trung tâm trong việc lưu trữ tài liệu, quản lý cây thư mục phân cấp, theo dõi đa phiên bản và kiểm soát phân quyền đa cấp nghiêm ngặt, gồm 8 thực thể:\n"
           "1. Bảng Thu_muc (Cây thư mục phân cấp):\n"
           "   - ma_thu_muc (BIGINT, PK): Mã định danh thư mục.\n"
           "   - ma_cty (BIGINT, FK -> Cong_ty.ma_cty): Doanh nghiệp sở hữu thư mục.\n"
           "   - thuoc_phong_ban (BIGINT, FK -> Phong_ban.ma_phong_ban): Phòng ban quản lý thư mục.\n"
           "   - thu_muc_cha (BIGINT, FK -> Thu_muc.ma_thu_muc): Khóa ngoại tự tham chiếu chính bảng Thu_muc, cho phép tổ chức cây thư mục lồng nhau không giới hạn cấp độ.\n"
           "   - ten_thu_muc (VARCHAR(255)): Tên hiển thị của thư mục.\n"
           "   - nguoi_tao (BIGINT, FK -> Nguoi_dung.ma_nv): Nhân viên thực hiện tạo thư mục.\n"
           "   - thoi_gian_tao_thu_muc (TIMESTAMP): Thời điểm tạo thư mục.\n"
           "\n"
           "2. Bảng Tep (Tệp tài liệu nội bộ):\n"
           "   - ma_tep (BIGINT, PK): Mã định danh tài liệu.\n"
           "   - ma_cty (BIGINT, FK -> Cong_ty.ma_cty): Doanh nghiệp sở hữu tệp.\n"
           "   - ma_thu_muc (BIGINT, FK -> Thu_muc.ma_thu_muc): Thư mục chứa tệp.\n"
           "   - ma_phong_ban (BIGINT, FK -> Phong_ban.ma_phong_ban): Phòng ban sở hữu tài liệu.\n"
           "   - nguoi_tai_len (BIGINT, FK -> Nguoi_dung.ma_nv): Người thực hiện tải tệp lên.\n"
           "   - ten_tep (VARCHAR(500)): Tên gốc của tệp tài liệu.\n"
           "   - dinh_dang_tep (VARCHAR(20)): Đuôi định dạng (pdf, docx, txt).\n"
           "   - kich_thuoc (BIGINT): Dung lượng tệp tin tính theo byte.\n"
           "   - pham_vi (VARCHAR(20)): Phạm vi truy cập (CA_NHAN, PHONG_BAN, DOANH_NGHIEP).\n"
           "   - ma_phien_ban (BIGINT, FK -> Phien_ban_tep.ma_phien_ban): Khóa ngoại chỉ định phiên bản đang có hiệu lực hiện hành.\n"
           "   - nhung_ai (BOOLEAN): Cờ đánh dấu tài liệu đã được hoàn tất lập chỉ mục vector cho trợ lý AI hay chưa.\n"
           "   - cl_url (TEXT): Đường dẫn định danh tệp tin trên kho lưu trữ đám mây.\n"
           "   - anh_thu_nho_url (TEXT): Đường dẫn hình ảnh thu nhỏ phục vụ xem trước (thumbnail).\n"
           "   - thoi_gian_tai (TIMESTAMP): Thời điểm tải tệp lên hệ thống.\n"
           "   - thoi_gian_cap_nhat (TIMESTAMP): Thời điểm cập nhật tệp gần nhất.\n"
           "   - thoi_gian_xoa (TIMESTAMP): Thời điểm chuyển vào thùng rác (hỗ trợ xóa mềm).\n"
           "\n"
           "3. Bảng Phien_ban_tep (Lịch sử các phiên bản tệp):\n"
           "   - ma_phien_ban (BIGINT, PK): Mã định danh phiên bản.\n"
           "   - ma_tep (BIGINT, FK -> Tep.ma_tep): Xác định phiên bản thuộc tệp nào (quan hệ 1-N: một tệp có nhiều phiên bản theo thời gian).\n"
           "   - so_thu_tu (INTEGER): Số thứ tự phiên bản (1, 2, 3...).\n"
           "   - AWS_url (TEXT): Đường dẫn lưu trữ tệp vật lý tương ứng trên kho S3.\n"
           "   - dung_luong (BIGINT): Kích thước tệp của phiên bản tương ứng.\n"
           "   - ng_tai (BIGINT, FK -> Nguoi_dung.ma_nv): Người tải lên phiên bản này.\n"
           "   - thoi_gian_tai (TIMESTAMP): Thời điểm tải lên phiên bản.\n"
           "\n"
           "4. Bảng Loai_tep (Danh mục nhãn tài liệu):\n"
           "   - ma_tep (BIGINT, PK, FK -> Tep.ma_tep): Mã tài liệu.\n"
           "   - loai (VARCHAR(100), PK): Tên thẻ/nhãn phân loại tài liệu (Tagging).\n"
           "\n"
           "5. Bảng Tep_danh_dau_sao (Tài liệu yêu thích):\n"
           "   - ma_danh_dau (BIGINT, PK): Mã định danh đánh dấu.\n"
           "   - ma_nv (BIGINT, FK -> Nguoi_dung.ma_nv): Nhân viên thực hiện đánh dấu sao.\n"
           "   - ma_tep (BIGINT, FK -> Tep.ma_tep): Tài liệu được đánh dấu yêu thích.\n"
           "   - thoi_gian_danh_dau (TIMESTAMP): Thời điểm thao tác đánh dấu sao.\n"
           "\n"
           "6. Bảng Quyen_thu_muc (Phân quyền trên thư mục):\n"
           "   - ma_quyen_thu_muc (BIGINT, PK): Mã định danh phân quyền thư mục.\n"
           "   - ma_thu_muc (BIGINT, FK -> Thu_muc.ma_thu_muc): Thư mục được áp dụng phân quyền.\n"
           "   - loai_doi_tuong (VARCHAR(20)): Đối tượng thụ hưởng quyền (NGUOI_DUNG hoặc PHONG_BAN).\n"
           "   - ma_nv (BIGINT, FK -> Nguoi_dung.ma_nv): Nhân viên cụ thể được gán quyền.\n"
           "   - muc_do (VARCHAR(20)): Mức quyền áp dụng (XEM, CHINH_SUA, QUAN_TRI).\n"
           "   - thoi_gian_phan_quyen (TIMESTAMP): Thời điểm thiết lập quyền.\n"
           "\n"
           "7. Bảng Quyen_tep (Phân quyền trên tệp tài liệu):\n"
           "   - ma_quyen (BIGINT, PK): Mã định danh quyền trên tệp.\n"
           "   - ma_tep (BIGINT, FK -> Tep.ma_tep): Tài liệu được phân quyền.\n"
           "   - phong_ban (BIGINT, FK -> Phong_ban.ma_phong_ban): Phòng ban được cấp quyền truy cập.\n"
           "   - ma_nv (BIGINT, FK -> Nguoi_dung.ma_nv): Nhân viên được cấp quyền truy cập trực tiếp.\n"
           "   - muc (VARCHAR(20)): Cấp độ quyền truy cập (XEM, CHINH_SUA, QUAN_TRI).\n"
           "   - ma_nguoi_cap (BIGINT, FK -> Nguoi_dung.ma_nv): Người thực hiện cấp quyền.\n"
           "   - thoi_diem_tao (TIMESTAMP): Thời điểm thiết lập phân quyền.\n"
           "\n"
           "8. Bảng Yeu_cau_cap_quyen (Phiếu yêu cầu cấp quyền tài liệu/thư mục):\n"
           "   - ma_yeu_cau (BIGINT, PK): Mã định danh phiếu yêu cầu.\n"
           "   - ma_thu_muc (BIGINT, FK -> Thu_muc.ma_thu_muc): Thư mục xin cấp quyền (nếu xin quyền thư mục).\n"
           "   - ten_thu_muc (VARCHAR(20)): Tên thư mục liên quan.\n"
           "   - ma_tep (BIGINT, FK -> Tep.ma_tep): Tệp tài liệu xin cấp quyền (nếu xin quyền tệp).\n"
           "   - ten_tep (VARCHAR(255)): Tên tệp liên quan.\n"
           "   - ng_yeu_cau (BIGINT, FK -> Nguoi_dung.ma_nv): Nhân viên gửi yêu cầu.\n"
           "   - ng_duyet (BIGINT, FK -> Nguoi_dung.ma_nv): Người có thẩm quyền tiếp nhận và phê duyệt.\n"
           "   - ly_do (TEXT): Lý do công việc cần truy cập tài liệu.\n"
           "   - muc_yeu_cau (VARCHAR(20)): Mức quyền xin cấp (XEM, CHINH_SUA, QUAN_TRI).\n"
           "   - trang_thai (VARCHAR(20)): Trạng thái xử lý (CHO_DUYET, DA_PHE_DUYET, TU_CHOI).\n"
           "   - thoi_gian_yeu_cau (TIMESTAMP): Thời điểm gửi yêu cầu.\n"
           "   - thoi_gian_phan_hoi (TIMESTAMP): Thời điểm cấp quản lý đưa ra quyết định duyệt/từ chối.")

# Subsystem 3: Tri thức & Vector AI
p_sub3 = doc.add_paragraph()
set_p_text(p_sub3, "5.2.3. Phân hệ Phân đoạn Tri thức và Chỉ mục Ngữ nghĩa (AI Pipeline)", bold=True, size_pt=12.5, space_after=4)

p_sub3_desc = doc.add_paragraph()
set_p_text(p_sub3_desc,
           "Phân hệ chịu trách nhiệm lưu trữ kết quả phân tách văn bản từ tài liệu gốc, làm dữ liệu đầu vào cho quy trình tìm kiếm ngữ nghĩa và trích dẫn nguồn của trợ lý thông minh:\n"
           "1. Bảng Cac_phan_tai_lieu (Các phân đoạn tài liệu - Text Chunks):\n"
           "   - ma_phan (BIGINT, PK): Mã định danh phân đoạn văn bản.\n"
           "   - ma_tep (BIGINT, FK -> Tep.ma_tep): Tài liệu gốc chứa phân đoạn này (quan hệ 1-N: một tài liệu được chia thành nhiều đoạn văn bản).\n"
           "   - stt_phan (INTEGER): Số thứ tự của đoạn văn bản trong toàn bộ tài liệu.\n"
           "   - ma_dinh_danh (VARCHAR(64)): UUID liên kết định danh sang vector embedding tương ứng lưu trong Vector Database.\n"
           "   - noi_dung (TEXT): Nội dung văn bản thuần túy của phân đoạn.\n"
           "   - so_luong_token (INTEGER): Số lượng token của đoạn văn bản phục vụ quản lý context window của mô hình ngôn ngữ lớn (LLM).\n"
           "   - trang (INTEGER): Số trang cụ thể trong tài liệu gốc nơi chứa đoạn văn bản này, dùng để hiển thị căn cứ dẫn chứng chuẩn xác khi trợ lý AI trả lời.\n"
           "   - ma_bam (VARCHAR(64)): Mã băm SHA-256 của nội dung đoạn văn bản nhằm đối soát tính toàn vẹn và loại bỏ trùng lặp khi tài liệu được cập nhật lại.\n"
           "   - trang_thai (BOOLEAN): Trạng thái kích hoạt phục vụ tra cứu của phân đoạn.\n"
           "   - ngay_bam (TIMESTAMP): Thời điểm thực hiện trích xuất và băm dữ liệu.")

# Subsystem 4: Tương tác Hội thoại & Nhật ký
p_sub4 = doc.add_paragraph()
set_p_text(p_sub4, "5.2.4. Phân hệ Tương tác Hội thoại, Đánh giá và Nhật ký Hoạt động", bold=True, size_pt=12.5, space_after=4)

p_sub4_desc = doc.add_paragraph()
set_p_text(p_sub4_desc,
           "Phân hệ quản lý toàn bộ các phiên trò chuyện giữa nhân viên và trợ lý AI, thu thập đánh giá chất lượng câu trả lời và ghi nhận nhật ký kiểm toán toàn diện cho hệ thống, gồm 4 thực thể:\n"
           "1. Bảng Phien_chat (Phiên hội thoại):\n"
           "   - ma_phien_chat (BIGINT, PK): Mã định danh phiên hội thoại.\n"
           "   - ma_nv (BIGINT, FK -> Nguoi_dung.ma_nv): Nhân viên sở hữu phiên hội thoại (quan hệ 1-N: một nhân viên có nhiều phiên hỏi đáp theo từng chủ đề).\n"
           "   - tieu_de (VARCHAR(255)): Tiêu đề gợi nhớ của phiên trò chuyện do người dùng đặt hoặc hệ thống tự động tóm tắt từ câu hỏi đầu tiên.\n"
           "   - thoi_gian_bat_dau (TIMESTAMP): Thời điểm khởi tạo cuộc trò chuyện.\n"
           "\n"
           "2. Bảng Tin_nhan (Nội dung trao đổi trong phiên):\n"
           "   - ma_tn (BIGINT, PK): Mã định danh tin nhắn.\n"
           "   - ma_phien_chat (BIGINT, FK -> Phien_chat.ma_phien_chat): Phiên hội thoại chứa tin nhắn (quan hệ 1-N: một phiên chat chứa nhiều tin nhắn).\n"
           "   - vai_tro (VARCHAR(20)): Người gửi tin nhắn, gồm: USER (người dùng đặt câu hỏi) hoặc ASSISTANT (trợ lý AI phản hồi).\n"
           "   - noi_dung (TEXT): Toàn văn nội dung câu hỏi hoặc câu trả lời tổng hợp kèm trích dẫn.\n"
           "   - bi_chan (BOOLEAN): Cờ đánh dấu nếu câu hỏi hoặc phản hồi bị kiểm duyệt do vi phạm chính sách bảo mật dữ liệu.\n"
           "   - thoi_gian_gui (TIMESTAMP): Thời điểm tin nhắn được gửi đi.\n"
           "\n"
           "3. Bảng Danh_gia_tin_nhan (Đánh giá chất lượng câu trả lời):\n"
           "   - ma_danh_gia (BIGINT, PK): Mã định danh bản ghi đánh giá.\n"
           "   - ma_tn (BIGINT, FK -> Tin_nhan.ma_tn): Tin nhắn của trợ lý được người dùng đánh giá.\n"
           "   - ma_nv (BIGINT, FK -> Nguoi_dung.ma_nv): Nhân viên thực hiện đánh giá.\n"
           "   - danh_gia (VARCHAR(10)): Mức độ hài lòng (LIKE hoặc DISLIKE).\n"
           "   - ly_do (VARCHAR(50)): Lý do chính nếu không hài lòng (ví dụ: SAI_THONG_TIN, KHONG_DU_CHI_TIET, TRICH_DAN_SAI).\n"
           "   - y_kien_cu_the (TEXT): Nhận xét chi tiết từ người dùng.\n"
           "   - mong_muon (TEXT): Góp ý của người dùng về nội dung mong muốn nhận được.\n"
           "   - thoi_gian_phan_hoi (TIMESTAMP): Thời điểm gửi đánh giá.\n"
           "\n"
           "4. Bảng Nhat_ky_hoat_dong (Nhật ký kiểm toán hệ thống - Audit Trail):\n"
           "   - ma_nhat_ky (BIGINT, PK): Mã định danh bản ghi nhật ký.\n"
           "   - ma_cty (BIGINT, FK -> Cong_ty.ma_cty): Doanh nghiệp nơi sự kiện diễn ra.\n"
           "   - ma_nv (BIGINT, FK -> Nguoi_dung.ma_nv): Người thực hiện hành động.\n"
           "   - hanh_dong (VARCHAR(20)): Loại hành động được ghi nhận (DANG_NHAP, DANG_XUAT, TAI_LEN, XOA_TEP, CAP_QUYEN, KHOI_PHUC, TRUY_VAN_AI).\n"
           "   - ma_tep (BIGINT, FK -> Tep.ma_tep): Tệp tài liệu liên quan đến hành động (nếu có).\n"
           "   - du_lieu_hanh_dong (JSONB): Chi tiết bổ sung về thao tác lưu dưới dạng JSON (địa chỉ IP, trình duyệt, nội dung thay đổi trước và sau khi thực hiện).\n"
           "   - thoi_gian_thuc_hien (TIMESTAMP): Dấu thời gian chính xác khi hành động diễn ra.")

print("Added new Section V content successfully!")

# -------------------------------------------------------------
# 4. UPDATE TOC AND FIGURE LIST
# -------------------------------------------------------------
# Let's inspect TOC and Figure list in doc.paragraphs:
# P15: V. THIẾT KẾ CƠ SỞ DỮ LIỆU MỨC LOGIC (ERD)	32
# P16: 5.1. ERD Phân hệ Tổ chức & Người dùng	32
# P17: 5.2. ERD Phân hệ Tài liệu & Phân quyền truy cập	33
# P18: 5.3. ERD Phân hệ Tri thức & Trợ lý thông minh	34
# P19: 5.4. ERD Phân hệ Hội thoại, Phản hồi & Nhật ký	35

set_p_text(doc.paragraphs[15], "V. THIẾT KẾ CƠ SỞ DỮ LIỆU MỨC LOGIC (ERD)\t32", bold=True, size_pt=12)
set_p_text(doc.paragraphs[16], "5.1. Sơ đồ quan hệ thực thể (ERD) tổng thể hệ thống\t32", bold=False, size_pt=12)
set_p_text(doc.paragraphs[17], "5.2. Đặc tả chi tiết các thực thể và ràng buộc dữ liệu\t33", bold=False, size_pt=12)
# Clear P18, P19 text (or delete them)
p19 = doc.paragraphs[19]
p19._p.getparent().remove(p19._p)
p18 = doc.paragraphs[18]
p18._p.getparent().remove(p18._p)

# Now check Figure catalog:
# Previously P57 was Hình 5.1, P58 Hình 5.2, P59 Hình 5.3, P60 Hình 5.4.
# Notice index shifted by -2 because we removed 2 paragraphs!
for idx, p in enumerate(doc.paragraphs):
    if "Hình 5.1:" in p.text and idx < 65:
        set_p_text(p, "Hình 5.1: Sơ đồ quan hệ thực thể (ERD) mức logic của hệ thống\t32", bold=False, size_pt=12)
    elif ("Hình 5.2:" in p.text or "Hình 5.3:" in p.text or "Hình 5.4:" in p.text) and idx < 65:
        p._p.getparent().remove(p._p)

print("Updated Table of Contents and Figure Catalog successfully!")

# Save back to DOCX
doc.save(DOCX_PATH)
print("File successfully saved to:", DOCX_PATH)
